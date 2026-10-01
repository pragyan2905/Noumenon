import fitz
from docx import Document
from docx.shared import Pt
from pathlib import Path
import tempfile

def test():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = 0
    section.bottom_margin = 0
    section.left_margin = 0
    section.right_margin = 0
    section.header_distance = 0
    section.footer_distance = 0
    
    pdf_path = "/Users/pragyan2905/file_converter/2309.07597v2.pdf"
    pdf_document = fitz.open(pdf_path)
    
    with tempfile.TemporaryDirectory() as temp_dir:
        for page_num in range(min(3, len(pdf_document))):
            page = pdf_document[page_num]
            
            if page_num > 0:
                doc.add_page_break()
                
            section.page_width = Pt(page.rect.width)
            section.page_height = Pt(page.rect.height)
            
            zoom = 150 / 72
            mat = fitz.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            
            img_path = Path(temp_dir) / f"page_{page_num}.png"
            pix.save(str(img_path))
            
            if page_num == 0:
                p = doc.paragraphs[0]
            else:
                p = doc.add_paragraph()
                
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            
            r = p.add_run()
            r.add_picture(str(img_path), width=Pt(page.rect.width), height=Pt(page.rect.height - 25))
            
    output = "test_output.docx"
    doc.save(output)
    print("Saved to", output)

if __name__ == "__main__":
    test()
