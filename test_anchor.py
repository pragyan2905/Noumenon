from docx import Document
from docx.shared import Pt
import fitz
from docx.oxml import parse_xml
import copy

def make_floating(inline_shape):
    # inline_shape._inline is the <wp:inline> element
    inline = inline_shape._inline
    
    # We want to replace <wp:inline> with <wp:anchor>
    # We can parse a template <wp:anchor> and copy children
    anchor_xml = """
    <wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
               xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
               xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"
               xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
               behindDoc="1" locked="0" layoutInCell="0" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page">
            <wp:posOffset>0</wp:posOffset>
        </wp:positionH>
        <wp:positionV relativeFrom="page">
            <wp:posOffset>0</wp:posOffset>
        </wp:positionV>
        <wp:extent cx="{cx}" cy="{cy}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:wrapNone/>
        <wp:docPr id="{id}" name="Picture {id}"/>
        <wp:cNvGraphicFramePr>
            <a:graphicFrameLocks noChangeAspect="1"/>
        </wp:cNvGraphicFramePr>
    </wp:anchor>
    """
    extent = inline.find('.//wp:extent', namespaces=inline.nsmap)
    cx = extent.get('cx')
    cy = extent.get('cy')
    docpr = inline.find('.//wp:docPr', namespaces=inline.nsmap)
    doc_id = docpr.get('id')
    
    anchor = parse_xml(anchor_xml.format(cx=cx, cy=cy, id=doc_id))
    
    # move graphic
    graphic = inline.find('.//a:graphic', namespaces=inline.nsmap)
    anchor.append(copy.deepcopy(graphic))
    
    # replace inline with anchor
    drawing = inline.getparent()
    drawing.replace(inline, anchor)

doc = Document()
doc.sections[0].page_width = Pt(595)
doc.sections[0].page_height = Pt(842)
doc.sections[0].top_margin = Pt(0)
doc.sections[0].bottom_margin = Pt(0)
doc.sections[0].left_margin = Pt(0)
doc.sections[0].right_margin = Pt(0)

# Create a red box image
import PIL.Image
img = PIL.Image.new('RGB', (595, 842), color = 'red')
img.save("red.jpg")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run()
shape = r.add_picture("red.jpg", width=Pt(595), height=Pt(842))
make_floating(shape)

# add second page
doc.add_page_break()
p2 = doc.add_paragraph()
r2 = p2.add_run()
shape2 = r2.add_picture("red.jpg", width=Pt(595), height=Pt(842))
make_floating(shape2)

doc.save("test_anchor.docx")
print("Saved test_anchor.docx")
