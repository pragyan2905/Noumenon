import time
import tempfile
import shutil
from pathlib import Path

try:
    import markdown
    HAS_MARKDOWN = True
except ImportError:
    HAS_MARKDOWN = False

try:
    from bs4 import BeautifulSoup
    HAS_BS4 = True
except ImportError:
    HAS_BS4 = False

try:
    from xhtml2pdf import pisa
    import docx
    HAS_PDF_DOCX = True
except ImportError:
    HAS_PDF_DOCX = False

from src.core.interfaces.engine import ConversionEngine
from src.core.models.conversion import ConversionRequest, ConversionResult

class TextEngine(ConversionEngine):
    """
    Handles robust text conversions: txt, md, html.
    """
    
    def __init__(self):
        self.formats = {'txt', 'md', 'html'}
        if HAS_PDF_DOCX:
            self.formats.update({'pdf', 'docx'})

    def can_convert(self, source_format: str, target_format: str) -> bool:
        return source_format in self.formats and target_format in self.formats

    def get_quality_score(self) -> str:
        return 'A'

    def convert(self, request: ConversionRequest) -> ConversionResult:
        source_ext = request.input_path.suffix.lower().lstrip('.')
        target_ext = request.output_format.lower().lstrip('.')
        
        start_time = time.time()
        
        try:
            if not request.input_path.exists():
                return ConversionResult(success=False, error_message=f"Input file not found: {request.input_path}")
            
            # Read input
            try:
                with open(request.input_path, 'r', encoding='utf-8') as f:
                    text_content = f.read()
            except Exception as e:
                return ConversionResult(success=False, error_message=f"Failed to read {source_ext.upper()}: {str(e)}")
            
            # Convert
            output_content = text_content
            
            if source_ext == 'md' and target_ext == 'html':
                if HAS_MARKDOWN:
                    output_content = markdown.markdown(text_content, extensions=['tables', 'fenced_code'])
                else:
                    return ConversionResult(success=False, error_message="Markdown library not installed for md->html conversion.")
            
            elif source_ext == 'html' and target_ext == 'md':
                # Basic conversion, ideally we'd use html2text or markdownify, but keeping dependencies light
                if HAS_BS4:
                    soup = BeautifulSoup(text_content, 'html.parser')
                    # Very rudimentary extraction
                    output_content = soup.get_text(separator='\n\n')
                else:
                    # Fallback to dumb tags stripping
                    import re
                    output_content = re.sub(r'<[^>]+>', '', text_content)
                    
            elif source_ext == 'html' and target_ext == 'txt':
                if HAS_BS4:
                    soup = BeautifulSoup(text_content, 'html.parser')
                    output_content = soup.get_text(separator='\n')
                else:
                    import re
                    output_content = re.sub(r'<[^>]+>', '', text_content)
                    
            elif source_ext == 'txt' and target_ext == 'html':
                # Escape and wrap in pre
                import html
                escaped = html.escape(text_content)
                output_content = f"<html><body><pre>{escaped}</pre></body></html>"
                
            elif target_ext in ('pdf', 'docx') and HAS_PDF_DOCX:
                # Convert source to HTML first if it isn't
                html_source = text_content
                if source_ext == 'md' and HAS_MARKDOWN:
                    html_source = markdown.markdown(text_content, extensions=['tables', 'fenced_code'])
                elif source_ext == 'txt':
                    import html
                    html_source = f"<html><body><pre>{html.escape(text_content)}</pre></body></html>"
                    
                # Handle PDF
                if target_ext == 'pdf':
                    styled_html = f"<html><head><style>body{{font-family: sans-serif; line-height: 1.5; padding: 1in;}}</style></head><body>{html_source}</body></html>"
                    pdf_out_path = request.output_directory / f"{request.input_path.stem}.pdf"
                    with open(pdf_out_path, "wb") as pdf_out:
                        pisa_status = pisa.CreatePDF(styled_html, dest=pdf_out)
                    if pisa_status.err:
                        raise ValueError("Failed to generate PDF from markup.")
                    return ConversionResult(success=True, output_path=pdf_out_path, engine_used=self.__class__.__name__, duration_seconds=time.time()-start_time)
                
                # Handle DOCX (Dumb HTML text extraction for basic support)
                elif target_ext == 'docx':
                    from docx import Document
                    doc = Document()
                    
                    if HAS_BS4:
                        soup = BeautifulSoup(html_source, 'html.parser')
                        for p in soup.find_all(['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li']):
                            if p.name.startswith('h'):
                                doc.add_heading(p.get_text(), level=int(p.name[1]))
                            elif p.name == 'li':
                                doc.add_paragraph(p.get_text(), style='List Bullet')
                            else:
                                doc.add_paragraph(p.get_text())
                        if not soup.find(['p', 'h1', 'li']): # Fallback
                             doc.add_paragraph(soup.get_text())
                    else:
                        import re
                        clean_text = re.sub(r'<[^>]+>', '', html_source)
                        doc.add_paragraph(clean_text)
                        
                    docx_out_path = request.output_directory / f"{request.input_path.stem}.docx"
                    doc.save(docx_out_path)
                    return ConversionResult(success=True, output_path=docx_out_path, engine_used=self.__class__.__name__, duration_seconds=time.time()-start_time)
                
            # Write temp output
            fd, temp_out_path = tempfile.mkstemp(suffix=f".{target_ext}")
            import os
            os.close(fd)
            temp_out = Path(temp_out_path)
            
            try:
                with open(temp_out, 'w', encoding='utf-8') as f:
                    f.write(output_content)
                
                # Verify generated output by checking size
                if temp_out.stat().st_size == 0 and len(output_content) > 0:
                    raise ValueError("Generated output file is unexpectedly empty.")
                
                output_path = request.output_directory / f"{request.input_path.stem}.{target_ext}"
                shutil.move(str(temp_out), str(output_path))
                
            except Exception as e:
                if temp_out.exists():
                    temp_out.unlink()
                raise e
                
            return ConversionResult(
                success=True,
                output_path=output_path,
                engine_used=self.__class__.__name__,
                duration_seconds=time.time() - start_time
            )
            
        except Exception as e:
            return ConversionResult(
                success=False,
                error_message=f"Text conversion failed: {str(e)}",
                duration_seconds=time.time() - start_time
            )
